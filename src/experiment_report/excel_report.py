"""정리 결과를 엑셀 보고서(표 + 차트)로 저장."""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .summarizer import SummaryResult

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REF_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(italic=True, size=9, color="666666")


def _style_header_row(ws: Worksheet, row: int, col_start: int, col_end: int) -> None:
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _autofit(ws: Worksheet, n_cols: int, widths: list[int] | None = None) -> None:
    for i in range(1, n_cols + 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths[i - 1] if widths else 16


def _write_organized_sheet(wb: Workbook, summary: SummaryResult) -> None:
    ws = wb.create_sheet("정리표")
    ws["A1"] = "실험 Raw Data 정리표"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"기준 샘플: {summary.reference}   |   생성 시각: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = SUBTITLE_FONT

    header_row = 4
    ws.cell(row=header_row, column=1, value="측정 항목")
    for j, sample in enumerate(summary.organized.columns, start=2):
        label = f"{sample} (기준)" if sample == summary.reference else sample
        ws.cell(row=header_row, column=j, value=label)
    _style_header_row(ws, header_row, 1, len(summary.organized.columns) + 1)

    for i, item in enumerate(summary.organized.index, start=header_row + 1):
        ws.cell(row=i, column=1, value=item).border = BORDER
        for j, sample in enumerate(summary.organized.columns, start=2):
            val = summary.organized.loc[item, sample]
            cell = ws.cell(row=i, column=j, value=None if val != val else round(float(val), 4))
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if sample == summary.reference:
                cell.fill = REF_FILL

    # 반복 측정 횟수(n) 표
    n_header_row = header_row + len(summary.organized.index) + 3
    ws.cell(row=n_header_row - 1, column=1, value="반복 측정 횟수 (n)").font = Font(bold=True)
    ws.cell(row=n_header_row, column=1, value="측정 항목")
    for j, sample in enumerate(summary.counts.columns, start=2):
        ws.cell(row=n_header_row, column=j, value=sample)
    _style_header_row(ws, n_header_row, 1, len(summary.counts.columns) + 1)
    for i, item in enumerate(summary.counts.index, start=n_header_row + 1):
        ws.cell(row=i, column=1, value=item).border = BORDER
        for j, sample in enumerate(summary.counts.columns, start=2):
            n = summary.counts.loc[item, sample]
            cell = ws.cell(row=i, column=j, value=int(n) if n == n else 0)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")

    _autofit(ws, len(summary.organized.columns) + 1)
    ws.freeze_panes = "B5"


def _write_comparison_sheet(wb: Workbook, summary: SummaryResult) -> None:
    ws = wb.create_sheet("기준대비 비교")
    ws["A1"] = f"기준 샘플({summary.reference}) 대비 비교"
    ws["A1"].font = TITLE_FONT

    if not summary.compare_samples:
        ws["A3"] = "비교할 샘플이 지정되지 않았습니다."
        return

    sub_cols = ["값", "기준 대비 차이", "기준 대비 차이(%)"]
    top_row, sub_row = 3, 4
    ws.cell(row=top_row, column=1, value="측정 항목")
    ws.merge_cells(start_row=top_row, start_column=1, end_row=sub_row, end_column=1)
    ws.cell(row=sub_row, column=1)

    col = 2
    pct_columns: list[int] = []
    for sample in summary.compare_samples:
        ws.cell(row=top_row, column=col, value=sample)
        ws.merge_cells(start_row=top_row, start_column=col, end_row=top_row, end_column=col + 2)
        for k, sub in enumerate(sub_cols):
            ws.cell(row=sub_row, column=col + k, value=sub)
        pct_columns.append(col + 2)
        col += 3
    last_col = col - 1
    _style_header_row(ws, top_row, 1, last_col)
    _style_header_row(ws, sub_row, 1, last_col)

    for i, item in enumerate(summary.item_labels, start=sub_row + 1):
        ws.cell(row=i, column=1, value=item).border = BORDER
        col = 2
        for sample in summary.compare_samples:
            df = summary.comparison[sample]
            row_vals = df.loc[item]
            for k, sub in enumerate(sub_cols):
                v = row_vals[sub]
                v = None if v != v else round(float(v), 4)
                cell = ws.cell(row=i, column=col + k, value=v)
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center")
            col += 3

    first_data_row = sub_row + 1
    last_data_row = sub_row + len(summary.item_labels)
    for pc in pct_columns:
        letter = get_column_letter(pc)
        rule = ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        )
        ws.conditional_formatting.add(f"{letter}{first_data_row}:{letter}{last_data_row}", rule)

    _autofit(ws, last_col)
    ws.freeze_panes = "B5"


def _write_chart_sheet(wb: Workbook, summary: SummaryResult) -> None:
    ws = wb.create_sheet("차트")
    ws["A1"] = "항목별 비교 차트"
    ws["A1"].font = TITLE_FONT

    data_start_row = 3
    n_samples = len(summary.organized.columns)
    ws.cell(row=data_start_row, column=1, value="측정 항목")
    for j, sample in enumerate(summary.organized.columns, start=2):
        ws.cell(row=data_start_row, column=j, value=sample)
    _style_header_row(ws, data_start_row, 1, n_samples + 1)

    for i, item in enumerate(summary.organized.index, start=data_start_row + 1):
        ws.cell(row=i, column=1, value=item).border = BORDER
        for j, sample in enumerate(summary.organized.columns, start=2):
            val = summary.organized.loc[item, sample]
            cell = ws.cell(row=i, column=j, value=None if val != val else round(float(val), 4))
            cell.border = BORDER

    last_data_row = data_start_row + len(summary.organized.index)
    _autofit(ws, n_samples + 1)

    chart_anchor_row = last_data_row + 3
    for offset, item in enumerate(summary.organized.index):
        item_row = data_start_row + 1 + offset
        chart = BarChart()
        chart.type = "col"
        chart.title = item
        chart.y_axis.title = item
        chart.x_axis.title = "샘플"
        chart.style = 10
        chart.height = 7
        chart.width = 14

        cats = Reference(ws, min_col=2, max_col=n_samples + 1, min_row=data_start_row, max_row=data_start_row)
        data = Reference(ws, min_col=1, max_col=n_samples + 1, min_row=item_row, max_row=item_row)
        chart.add_data(data, titles_from_data=True, from_rows=True)
        chart.set_categories(cats)
        chart.legend = None

        anchor_col = 1 + (offset % 2) * 8
        anchor_row = chart_anchor_row + (offset // 2) * 16
        ws.add_chart(chart, f"{get_column_letter(anchor_col)}{anchor_row}")


def write_report(summary: SummaryResult, output_path: str) -> None:
    """정리표 / 기준대비 비교 / 차트 시트를 담은 엑셀 보고서를 생성한다."""
    wb = Workbook()
    wb.remove(wb.active)
    _write_organized_sheet(wb, summary)
    _write_comparison_sheet(wb, summary)
    _write_chart_sheet(wb, summary)
    wb.save(output_path)
