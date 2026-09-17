import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from experiment_report.data_loader import load_raw_data
from experiment_report.pipeline import RunConfig, run
from experiment_report.summarizer import build_summary

EXAMPLE_XLSX = Path(__file__).parent.parent / "data" / "raw_data_example.xlsx"


def test_load_raw_data_averages_replicates():
    result = load_raw_data(str(EXAMPLE_XLSX))
    assert set(result.item_codes) == {"TSC", "RTVM", "Mn", "Mw", "PDI"}
    assert list(result.averaged.index) == ["REF", "A-01", "A-02", "A-03"]
    assert result.averaged.loc["REF", "TSC"] == pytest.approx((45.2 + 45.4) / 2)
    assert result.counts.loc["REF", "TSC"] == 2


def test_load_raw_data_missing_sample_column_raises(tmp_path):
    import pandas as pd

    bad = tmp_path / "bad.xlsx"
    pd.DataFrame({"foo": [1, 2], "TSC(%)": [1.0, 2.0]}).to_excel(bad, index=False)
    with pytest.raises(ValueError, match="샘플명"):
        load_raw_data(str(bad))


def test_build_summary_computes_diff_vs_reference():
    result = load_raw_data(str(EXAMPLE_XLSX))
    summary = build_summary(result, reference="REF", samples=["A-01", "A-02"])
    assert summary.compare_samples == ["A-01", "A-02"]
    tsc_ref = summary.organized.loc["TSC (%)", "REF"]
    tsc_a01 = summary.organized.loc["TSC (%)", "A-01"]
    diff = summary.comparison["A-01"].loc["TSC (%)", "기준 대비 차이"]
    assert diff == pytest.approx(tsc_a01 - tsc_ref)


def test_build_summary_unknown_sample_raises():
    result = load_raw_data(str(EXAMPLE_XLSX))
    with pytest.raises(ValueError, match="찾을 수 없습니다"):
        build_summary(result, reference="REF", samples=["NOPE"])


def test_full_pipeline_writes_excel_with_charts(tmp_path):
    output = tmp_path / "report.xlsx"
    config = RunConfig(
        raw_data_path=str(EXAMPLE_XLSX),
        reference="REF",
        samples=["A-01", "A-02", "A-03"],
        output_path=str(output),
    )
    summary = run(config)
    assert output.exists()
    assert summary.reference == "REF"

    wb = load_workbook(output)
    assert wb.sheetnames == ["정리표", "기준대비 비교", "차트"]
    chart_ws = wb["차트"]
    assert len(chart_ws._charts) == len(summary.item_labels)
